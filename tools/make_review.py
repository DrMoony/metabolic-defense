#!/usr/bin/env python3
"""MI 검토본 생성: Word(가로 7열 + 문헌 부록) + Excel(검토의견 열). 실행: python3 tools/make_review.py v6"""
import json, datetime, sys, os
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ver=sys.argv[1] if len(sys.argv)>1 else 'vX'
D=os.path.expanduser("~/Library/CloudStorage/GoogleDrive-mftsky@gmail.com/내 드라이브/1_BI/Metabolic Defense")
today=datetime.date.today().strftime('%Y%m%d')
refs=json.load(open(f'{ROOT}/knowledge/references.json',encoding='utf-8'))
L=lambda n: json.load(open(f'{ROOT}/assets/{n}.json',encoding='utf-8'))
sets={'OB':('Clinical Obesity',L('quiz_obesity_ko'),L('quiz_obesity_en')),'MA':('MASLD / MASH',L('quiz_aasld_ko'),L('quiz_aasld_en'))}
used=sorted({k for v in sets.values() for r in v[1] for k in r.get('refs',[])})
acc={}
for v in sets.values():
    for r in v[1]:
        e=r.get('evidence') or {}; acc.setdefault(e.get('ref'),e.get('access'))
n_ab=sum(1 for v in sets.values() for r in v[1] if (r.get('evidence') or {}).get('access')=='abstract')
def ev(r):
    e=r.get('evidence') or {}
    if not e.get('quote'): return '(원문 구문 없음)'
    t=f"[{'직접' if e.get('support')=='direct' else '부분'} · {'전문' if e.get('access')=='fulltext' else '초록만 확인'}] {e.get('locator','')}\n“{e['quote'].strip()}”"
    if e.get('quote2'): t+=f"\n“{e['quote2'].strip()}”"
    if e.get('note') and e.get('support')!='direct': t+=f"\n※ {e['note']}"
    return t
doc=Document(); st=doc.styles['Normal']; st.font.name='Malgun Gothic'; st.font.size=Pt(9); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
for s in doc.sections:
    s.orientation=WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin=s.right_margin=Cm(1.2); s.top_margin=s.bottom_margin=Cm(1.2)
doc.add_heading(f'메타볼릭 디펜스 문제은행 — Medical Review 검토본 {ver}',0)
doc.add_paragraph(f'{datetime.date.today().isoformat()} · 총 {sum(len(v[1]) for v in sets.values())}문항 (Clinical Obesity {len(sets["OB"][1])} · MASLD/MASH {len(sets["MA"][1])}) · 참고문헌 {len(used)}종 · 전 문항 원문 근거 직접 인용')
doc.add_heading('검토 안내',1)
for t in ['대상: 학회 부스 라이트건 퀴즈(간호사·의료진). 정답 후 해설과 출처(APA, 화면에선 DOI 생략)가 표시됩니다.',
          '정책: 특정 약제(성분·브랜드·계열·임상시험) 문항 없음, 언브랜디드 질환 인지 콘텐츠, 쉬운 난이도 위주.',
          '근거: 문항마다 APA 7판 참고문헌과 그 문헌에서 글자 그대로 가져온 근거 구문을 붙였습니다.',
          f'근거 접근 범위: [전문]은 오픈액세스 전문·학회 원문 PDF·가이드라인 전문에서, [초록만 확인]은 유료 논문이라 PubMed 초록에서 인용했습니다({n_ab}문항). 초록 인용 문항은 사내 전문 접근으로 본문 대조를 부탁드립니다.',
          '검토 요청: ① 사실 정확성 ② 최신 가이드라인 일치 ③ 표현 적절성(낙인·과장) ④ 출처·근거 구문 적합성. 의견은 문항 ID 기준으로 남겨 주세요.']:
    doc.add_paragraph(t, style='List Bullet')
for key,(name,ko,en) in sets.items():
    doc.add_heading(f'{name} ({len(ko)}문항)',1)
    tbl=doc.add_table(rows=1, cols=7); tbl.style='Table Grid'
    for i,h in enumerate(['ID','난이도','문항','보기 (✓ 정답)','해설','참고문헌 (APA)','원문 근거 구문']):
        c=tbl.rows[0].cells[i]; c.text=h; c.paragraphs[0].runs[0].bold=True
    widths=[Cm(1.4),Cm(1.1),Cm(4.6),Cm(4.4),Cm(3.8),Cm(5.2),Cm(6.0)]
    for i,(r,e) in enumerate(zip(ko,en)):
        row=tbl.add_row().cells
        row[0].text=f'{key}-{i+1:03d}'; row[1].text=r['diff']; row[2].text=r['q']+'\n\n'+e['q']
        row[3].text='\n'.join(f"{'✓ ' if k==r['correct'] else '  '}{k+1}. {a}" for k,a in enumerate(r['a']))
        row[4].text=r.get('exp',''); row[5].text='\n\n'.join(r.get('src','').split(' / ')); row[6].text=ev(r)
        for k,c in enumerate(row):
            c.width=widths[k]
            for pp in c.paragraphs:
                for run in pp.runs: run.font.size=Pt(8)
    doc.add_paragraph()
# ---- 문헌별 문항 집계 ----
from collections import defaultdict
byref=defaultdict(list)
for key,(name,ko,en) in sets.items():
    for i,r in enumerate(ko):
        byref[(r.get('evidence') or {}).get('ref') or (r.get('refs') or ['?'])[0]].append(f'{key}-{i+1:03d}')
doc.add_heading('부록 A — 근거 문헌별 문항 집계',1)
t2=doc.add_table(rows=1, cols=4); t2.style='Table Grid'
for i,h in enumerate(['근거 문헌 (APA 앞부분)','접근','문항 수','문항 ID']):
    c=t2.rows[0].cells[i]; c.text=h; c.paragraphs[0].runs[0].bold=True
for k,ids in sorted(byref.items(), key=lambda x:-len(x[1])):
    row=t2.add_row().cells; row[0].text=refs.get(k,k)[:90]+('…' if len(refs.get(k,k))>90 else ''); row[1].text='전문' if acc.get(k)=='fulltext' else '초록'; row[2].text=str(len(ids)); row[3].text=', '.join(ids)
    for c in row:
        for pp in c.paragraphs:
            for run in pp.runs: run.font.size=Pt(8)
doc.add_paragraph()
doc.add_heading('부록 B — 참고문헌 (APA 7판, 사용된 문헌만; [전문]/[초록] = 근거 확인 범위)',1)
for k in sorted(used, key=lambda x: refs[x].lower()): doc.add_paragraph(f"[{'전문' if acc.get(k)=='fulltext' else '초록'}] {refs[k]}")
docx_path=f'{D}/[{today}]_MI리뷰_메타볼릭디펜스_문제은행_{ver}.docx'; doc.save(docx_path)
wb=Workbook(); ws0=wb.active; ws0.title='안내'
for line in [f'메타볼릭 디펜스 문제은행 {ver} — Medical Review 검토본', f'생성일 {datetime.date.today().isoformat()}', '시트: OB(Clinical Obesity), MA(MASLD/MASH), References', '근거 구문 = 참고문헌 원문에서 글자 그대로 복사. 근거 접근: 전문 / 초록만(사내 전문 대조 요청)', '검토 의견은 각 행의 "검토의견" 열에 적어 주세요.']:
    ws0.append([line])
for key,(name,ko,en) in sets.items():
    ws=wb.create_sheet(key)
    ws.append(['ID','난이도','문항(KO)','보기1','보기2','보기3','보기4','정답번호','해설(KO)','문항(EN)','정답(EN)','해설(EN)','참고문헌(APA)','근거 구문(원문)','근거 위치','근거 접근','비고','검토의견'])
    for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDEBF7')
    for i,(r,e) in enumerate(zip(ko,en)):
        v=r.get('evidence') or {}
        ws.append([f'{key}-{i+1:03d}',r['diff'],r['q'],*r['a'],r['correct']+1,r.get('exp',''),e['q'],e['a'][e['correct']],e.get('exp',''),'\n'.join(r.get('src','').split(' / ')),(v.get('quote','')+('\n'+v['quote2'] if v.get('quote2') else '')).strip(),v.get('locator',''),'전문' if v.get('access')=='fulltext' else '초록만',v.get('note','') if v.get('support')!='direct' else '',''])
    for col,w in zip('ABCDEFGHIJKLMNOPQR',[9,7,40,22,22,22,22,7,40,34,22,34,56,60,22,9,36,30]): ws.column_dimensions[col].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb2=wb.create_sheet('ByReference'); wb2.append(['근거 문헌 키','접근','문항 수','문항 ID','APA'])
for c in wb2[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDEBF7')
for k,ids in sorted(byref.items(), key=lambda x:-len(x[1])): wb2.append([k,'전문' if acc.get(k)=='fulltext' else '초록',len(ids),', '.join(ids),refs.get(k,k)])
for col,w in zip('ABCDE',[18,8,8,60,120]): wb2.column_dimensions[col].width=w
for row in wb2.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wr=wb.create_sheet('References'); wr.append(['키','접근','APA'])
for k in sorted(used, key=lambda x: refs[x].lower()): wr.append([k,'전문' if acc.get(k)=='fulltext' else '초록',refs[k]])
wr.column_dimensions['A'].width=18; wr.column_dimensions['C'].width=140
xlsx_path=f'{D}/[{today}]_MI리뷰_메타볼릭디펜스_문제은행_{ver}.xlsx'; wb.save(xlsx_path); print(docx_path); print(xlsx_path)
